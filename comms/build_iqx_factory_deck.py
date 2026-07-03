"""
Build "The IQX Factory" executive deck (clean, non-branded default design).

Source brief: comms/iqx-factory-exec-deck-brief.md
Output:       comms/IQX-Factory-Exec-Deck.pptx

Design goals (per brief Part 3):
- 4 slides, 16:9, executive audience, outcome-oriented.
- Rich but breathable: one headline, one core message, 3-4 short lines, one visual.
- Substance over generic; no internal mechanics; no invented metrics.

Run:  python comms/build_iqx_factory_deck.py
Deps: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
INK        = RGBColor(0x0F, 0x1B, 0x2D)   # near-black navy (primary text)
NAVY       = RGBColor(0x14, 0x2A, 0x4A)   # deep navy
ACCENT     = RGBColor(0x2E, 0x8B, 0xC0)   # confident blue
ACCENT_DK  = RGBColor(0x1B, 0x5E, 0x8A)   # darker blue
TEAL       = RGBColor(0x14, 0xB8, 0xA6)   # accent 2
SLATE      = RGBColor(0x5B, 0x6B, 0x7B)   # secondary text
MIST       = RGBColor(0xED, 0xF2, 0xF6)   # light panel
MIST2      = RGBColor(0xE0, 0xE8, 0xEF)   # slightly darker panel
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LINE       = RGBColor(0xCF, 0xDA, 0xE3)

HEAD_FONT = "Segoe UI Semibold"
BODY_FONT = "Segoe UI"
LIGHT_FONT = "Segoe UI Light"

# ---------------------------------------------------------------- setup
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, shadow=False,
         shape=MSO_SHAPE.RECTANGLE, round_amt=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shadow:
        _soft_shadow(sp)
    if round_amt is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = round_amt
        except Exception:
            pass
    return sp


def _soft_shadow(sp):
    spPr = sp._element.spPr
    effLst = spPr.makeelement(qn('a:effectLst'), {})
    shdw = spPr.makeelement(qn('a:outerShdw'), {
        'blurRad': '90000', 'dist': '38100', 'dir': '5400000', 'rotWithShape': '0'})
    clr = spPr.makeelement(qn('a:srgbClr'), {'val': '1B2A3A'})
    alpha = spPr.makeelement(qn('a:alpha'), {'val': '22000'})
    clr.append(alpha)
    shdw.append(clr)
    effLst.append(shdw)
    spPr.append(effLst)


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is a list of (txt, size, color,
    font, bold, spacing_pts) run-tuples. Simplify with single-run paragraphs."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, font, bold, tracking) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = font
            r.font.bold = bold
            if tracking:
                _set_tracking(r, tracking)
    return tb


def _set_tracking(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(int(pts * 100)))


def para(txt, size, color, font=BODY_FONT, bold=False, tracking=0, line_spacing=None):
    # line_spacing is accepted for call-site convenience; paragraph line spacing
    # is applied by text() via its own line_spacing argument.
    return [(txt, size, color, font, bold, tracking)]


def kicker(s, x, y, txt, color=ACCENT):
    text(s, x, y, Inches(8), Inches(0.35),
         [para(txt.upper(), 12.5, color, HEAD_FONT, True, tracking=2.2)])


def page_num(s, n):
    text(s, SW - Inches(0.9), SH - Inches(0.55), Inches(0.6), Inches(0.3),
         [para(str(n), 10, SLATE, BODY_FONT)], align=PP_ALIGN.RIGHT)


def footer(s):
    text(s, Inches(0.9), SH - Inches(0.55), Inches(6), Inches(0.3),
         [para("Verndale  ·  Data & Analytics", 9.5, SLATE, BODY_FONT, tracking=0.8)])


# ================================================================ SLIDE 1
# From bespoke to repeatable (the hook)
s = slide()
bg(s, WHITE)
# left accent band
rect(s, 0, 0, Inches(0.28), SH, fill=ACCENT)
# subtle top-right geometric accent
rect(s, SW - Inches(3.4), -Inches(1.2), Inches(4.6), Inches(4.6),
     fill=MIST, shape=MSO_SHAPE.OVAL)
rect(s, SW - Inches(2.1), -Inches(0.7), Inches(2.6), Inches(2.6),
     fill=MIST2, shape=MSO_SHAPE.OVAL)

kicker(s, Inches(0.9), Inches(0.85), "From bespoke to repeatable")

text(s, Inches(0.85), Inches(1.35), Inches(9.5), Inches(1.4),
     [para("The IQX Factory", 54, INK, HEAD_FONT, True)])

text(s, Inches(0.9), Inches(2.55), Inches(9.6), Inches(0.7),
     [para("Turning hand-built offerings into a repeatable engine",
           22, ACCENT_DK, LIGHT_FONT)])

# progression: craft -> engine -> many offerings
py = Inches(3.7)
pw, ph = Inches(3.15), Inches(1.55)
gap = Inches(0.55)
px = Inches(0.9)

stages = [
    ("Hand-built craft", "Expert, valuable — but slow and hard to scale", MIST, INK, SLATE),
    ("An AI-driven engine", "Verndale's methodology, encoded and repeatable", NAVY, WHITE, MIST),
    ("Many offerings", "Point it at an industry; a market-ready offering comes out", ACCENT, WHITE, MIST),
]
xcur = px
box_xs = []
for i, (title, sub, fill, tcol, scol) in enumerate(stages):
    b = rect(s, xcur, py, pw, ph, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             round_amt=0.08, shadow=True)
    text(s, xcur + Inches(0.25), py + Inches(0.22), pw - Inches(0.5), Inches(0.5),
         [para(title, 17, tcol, HEAD_FONT, True)])
    text(s, xcur + Inches(0.25), py + Inches(0.72), pw - Inches(0.5), Inches(0.75),
         [para(sub, 12.5, scol, BODY_FONT)], line_spacing=1.05)
    box_xs.append(xcur)
    xcur = xcur + pw + gap

# arrows between boxes
for i in range(len(stages) - 1):
    ax = box_xs[i] + pw + Inches(0.05)
    arr = rect(s, ax, py + ph/2 - Inches(0.16), gap - Inches(0.1), Inches(0.32),
               fill=ACCENT_DK, shape=MSO_SHAPE.CHEVRON)

# takeaway line
text(s, Inches(0.9), py + ph + Inches(0.5), Inches(11.2), Inches(0.6),
     [para("Verndale has proven it can build industry offerings — BankIQX is the first. "
           "The Factory lets us produce them repeatably, for any industry.",
           14, SLATE, BODY_FONT, line_spacing=1.15)])

footer(s)
page_num(s, 1)


# ================================================================ SLIDE 2
# One run, a market-ready offering (the substance)
s = slide()
bg(s, WHITE)
rect(s, 0, 0, SW, Inches(1.7), fill=NAVY)
rect(s, 0, Inches(1.7), SW, Inches(0.06), fill=ACCENT)

kicker(s, Inches(0.9), Inches(0.42), "A complete offering, not a document", color=TEAL)
text(s, Inches(0.85), Inches(0.82), Inches(11.5), Inches(0.8),
     [para("One run, a market-ready offering", 34, WHITE, HEAD_FONT, True)])

text(s, Inches(0.9), Inches(1.95), Inches(11.5), Inches(0.5),
     [para("Everything needed to take an industry offering to market — delivered as one package.",
           15, SLATE, BODY_FONT)])

tiles = [
    ("The commercial angle",
     "The validated wedge — where we win and why", ACCENT),
    ("The offering",
     "A use-case library and the customer-360 data model behind it", TEAL),
    ("The go-to-market",
     "Value case, positioning, and sales enablement", ACCENT_DK),
    ("The proof",
     "A working Sigma + Snowflake prototype prospects can see and click", NAVY),
]

# 2x2 grid
tw, th = Inches(5.65), Inches(1.95)
tgap = Inches(0.35)
tx0 = Inches(0.9)
ty0 = Inches(2.75)
positions = [(tx0, ty0), (tx0 + tw + tgap, ty0),
             (tx0, ty0 + th + tgap), (tx0 + tw + tgap, ty0 + th + tgap)]

for (title, body, accent), (tx, ty) in zip(tiles, positions):
    card = rect(s, tx, ty, tw, th, fill=MIST, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                round_amt=0.06, shadow=True)
    # accent edge
    rect(s, tx, ty, Inches(0.14), th, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    text(s, tx + Inches(0.45), ty + Inches(0.32), tw - Inches(0.8), Inches(0.6),
         [para(title, 20, INK, HEAD_FONT, True)])
    text(s, tx + Inches(0.45), ty + Inches(0.95), tw - Inches(0.8), Inches(0.85),
         [para(body, 14.5, SLATE, BODY_FONT, line_spacing=1.15)])

footer(s)
page_num(s, 2)


# ================================================================ SLIDE 3
# Why it matters (fast, credible, built to scale)
s = slide()
bg(s, WHITE)
rect(s, 0, 0, Inches(0.28), SH, fill=TEAL)

kicker(s, Inches(0.9), Inches(0.7), "Fast, credible, and built to scale")
text(s, Inches(0.85), Inches(1.1), Inches(11.5), Inches(0.8),
     [para("Why it matters", 38, INK, HEAD_FONT, True)])
text(s, Inches(0.9), Inches(1.95), Inches(11.2), Inches(0.5),
     [para("The Factory turns offering creation into a compounding, low-risk advantage.",
           15.5, SLATE, BODY_FONT)])

benefits = [
    ("Speed", "Enter new industries in a fraction of the usual time", ACCENT),
    ("Consistency", "The same rigor and standard every time, whoever runs it", TEAL),
    ("Credible by design", "Human-approved at key steps, evidence-backed claims, a technically honest prototype", ACCENT_DK),
    ("Scale", "One engine, many industries — each new offering easier than the last", NAVY),
]

bw, bh = Inches(2.78), Inches(2.35)
bgap = Inches(0.22)
bx0 = Inches(0.9)
by0 = Inches(2.75)
for i, (title, body, accent) in enumerate(benefits):
    bx = bx0 + i * (bw + bgap)
    rect(s, bx, by0, bw, bh, fill=WHITE, line=LINE, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_amt=0.05, shadow=True)
    # number chip
    rect(s, bx + Inches(0.3), by0 + Inches(0.3), Inches(0.55), Inches(0.55),
         fill=accent, shape=MSO_SHAPE.OVAL)
    text(s, bx + Inches(0.3), by0 + Inches(0.37), Inches(0.55), Inches(0.42),
         [para(str(i + 1), 18, WHITE, HEAD_FONT, True)], align=PP_ALIGN.CENTER)
    text(s, bx + Inches(0.3), by0 + Inches(1.05), bw - Inches(0.6), Inches(0.55),
         [para(title, 18, INK, HEAD_FONT, True)])
    text(s, bx + Inches(0.3), by0 + Inches(1.55), bw - Inches(0.6), Inches(0.75),
         [para(body, 12.5, SLATE, BODY_FONT, line_spacing=1.12)])

footer(s)
page_num(s, 3)


# ================================================================ SLIDE 4
# BankIQX today. A portfolio tomorrow. (the vision)
s = slide()
bg(s, NAVY)
rect(s, 0, 0, SW, SH, fill=NAVY)
# faint accent orb
rect(s, -Inches(1.5), SH - Inches(3.0), Inches(5.0), Inches(5.0),
     fill=ACCENT_DK, shape=MSO_SHAPE.OVAL)

kicker(s, Inches(0.9), Inches(0.85), "A growing family of IQX offerings", color=TEAL)
text(s, Inches(0.85), Inches(1.35), Inches(11.5), Inches(1.3),
     [para("BankIQX today. A portfolio tomorrow.", 44, WHITE, HEAD_FONT, True)])

text(s, Inches(0.9), Inches(2.55), Inches(11.0), Inches(0.6),
     [para("The same engine widens Verndale's moat with every new offering.",
           17, MIST, LIGHT_FONT)])

# roadmap band
rw, rh = Inches(2.65), Inches(1.85)
rgap = Inches(0.3)
rx0 = Inches(0.9)
ry0 = Inches(3.6)
roadmap = [
    ("BankIQX", "In market today", True),
    ("StudentIQX", "Next on the same engine", False),
    ("Industry 3", "Stood up faster", False),
    ("And beyond", "A widening library of reusable IP", False),
]
for i, (name, sub, filled) in enumerate(roadmap):
    rx = rx0 + i * (rw + rgap)
    if filled:
        rect(s, rx, ry0, rw, rh, fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             round_amt=0.07, shadow=True)
        nm_col, sub_col = WHITE, MIST
    else:
        c = rect(s, rx, ry0, rw, rh, fill=None, line=RGBColor(0x3C, 0x59, 0x7A),
                 line_w=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_amt=0.07)
        nm_col, sub_col = WHITE, SLATE
    text(s, rx + Inches(0.28), ry0 + Inches(0.45), rw - Inches(0.5), Inches(0.6),
         [para(name, 21, nm_col, HEAD_FONT, True)])
    text(s, rx + Inches(0.28), ry0 + Inches(1.05), rw - Inches(0.5), Inches(0.6),
         [para(sub, 13, sub_col, BODY_FONT, line_spacing=1.1)])

# closing line
text(s, Inches(0.9), Inches(5.95), Inches(11.4), Inches(0.7),
     [para("A compounding asset — not one-off work.",
           16, TEAL, HEAD_FONT, True)])

text(s, SW - Inches(3.4), SH - Inches(0.55), Inches(2.5), Inches(0.3),
     [para("Verndale  ·  Data & Analytics", 9.5, SLATE, BODY_FONT, tracking=0.8)],
     align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- save
import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "IQX-Factory-Exec-Deck.pptx")
prs.save(out)
print("Saved:", out)
